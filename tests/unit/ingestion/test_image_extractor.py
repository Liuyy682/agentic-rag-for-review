import io
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path


from agentic_rag.ingestion.image_extractor import (
    extract_images_from_pptx,
    extract_images_from_docx,
    clean_docx_broken_image_refs,
)


class TestPptxImageExtraction(unittest.TestCase):
    def test_extracts_picture_shapes_to_output_dir(self):
        """Images from PPTX picture shapes should be saved with MarkItDown-compatible names."""
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            pptx_path = temp_path / "test.pptx"
            output_dir = temp_path / "images" / "test"

            # Create a tiny valid PNG
            from PIL import Image
            png_path = temp_path / "tiny.png"
            Image.new("RGB", (100, 60), color="red").save(png_path, "PNG")

            # Build a PPTX with a picture shape
            prs = Presentation()
            slide_layout = prs.slide_layouts[6]  # blank
            slide = prs.slides.add_slide(slide_layout)
            slide.shapes.add_picture(
                str(png_path),
                Inches(1), Inches(1),
                Inches(2), Inches(1.5),
            )
            prs.save(str(pptx_path))

            extract_images_from_pptx(pptx_path, output_dir)

            # The shape default name is something like "Picture 1" → "Picture1.jpg"
            extracted = list(output_dir.iterdir())
            self.assertTrue(len(extracted) > 0, f"No images extracted to {output_dir}")
            self.assertTrue(
                any(f.suffix.lower() in (".jpg", ".jpeg", ".png") for f in extracted),
                f"Expected image files, got: {[f.name for f in extracted]}",
            )

    def test_output_dir_created_when_missing(self):
        """extract_images_from_pptx should create the output directory if it doesn't exist."""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            pptx_path = temp_path / "empty.pptx"
            output_dir = temp_path / "nonexistent" / "images"

            # Create a minimal PPTX with no pictures (just an empty presentation)
            from pptx import Presentation
            prs = Presentation()
            prs.slides.add_slide(prs.slide_layouts[6])  # blank slide, no pictures
            prs.save(str(pptx_path))

            extract_images_from_pptx(pptx_path, output_dir)

            # Directory should exist even if no images were extracted
            self.assertTrue(output_dir.exists())

    def test_pptx_with_no_pictures_is_noop(self):
        """A PPTX with no picture shapes should not produce any image files."""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            pptx_path = temp_path / "text_only.pptx"
            output_dir = temp_path / "images"
            output_dir.mkdir()

            from pptx import Presentation
            from pptx.util import Inches
            prs = Presentation()
            slide_layout = prs.slide_layouts[6]
            slide = prs.slides.add_slide(slide_layout)
            slide.shapes.add_textbox(
                Inches(1), Inches(1),
                Inches(3), Inches(1),
            ).text_frame.text = "Hello"
            prs.save(str(pptx_path))

            extract_images_from_pptx(pptx_path, output_dir)

            # Should not have created any image files
            image_files = [f for f in output_dir.iterdir() if f.suffix.lower() in (".jpg", ".jpeg", ".png")]
            self.assertEqual(len(image_files), 0)


class TestDocxImageExtraction(unittest.TestCase):
    def test_extracts_images_from_word_media(self):
        """Images under word/media/ in a DOCX ZIP should be extracted."""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            docx_path = temp_path / "test.docx"
            output_dir = temp_path / "images" / "test"

            # Build a minimal DOCX ZIP with an image in word/media/
            from PIL import Image
            img_bytes_io = io.BytesIO()
            Image.new("RGB", (80, 60), color="blue").save(img_bytes_io, "PNG")
            img_data = img_bytes_io.getvalue()

            with zipfile.ZipFile(docx_path, "w") as zf:
                zf.writestr("word/media/image1.png", img_data)
                zf.writestr("word/document.xml", "<document/>")
                zf.writestr("[Content_Types].xml", "<Types/>")

            hash_map = extract_images_from_docx(docx_path, output_dir)

            self.assertIn("image1.png", hash_map.values())
            extracted_path = output_dir / "image1.png"
            self.assertTrue(extracted_path.exists())
            self.assertEqual(extracted_path.read_bytes(), img_data)

    def test_docx_without_word_media_is_noop(self):
        """A DOCX without word/media/ should return an empty hash map."""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            docx_path = temp_path / "no_images.docx"
            output_dir = temp_path / "images"
            output_dir.mkdir()

            with zipfile.ZipFile(docx_path, "w") as zf:
                zf.writestr("word/document.xml", "<document/>")

            hash_map = extract_images_from_docx(docx_path, output_dir)
            self.assertEqual(hash_map, {})


class TestDocxDataUriCleanup(unittest.TestCase):
    def test_replaces_broken_data_uri_with_extracted_file(self):
        """Broken data:image refs should be replaced with extracted image filenames."""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            image_dir = temp_path / "images"
            image_dir.mkdir(parents=True)

            # Create extracted images
            (image_dir / "image1.png").write_bytes(b"fake-png-data")
            (image_dir / "image2.jpg").write_bytes(b"fake-jpg-data")

            markdown = (
                "Some text\n\n"
                "![alt](data:image/png;base64...)\n\n"
                "More text\n\n"
                "![](data:image/jpeg;base64...)\n\n"
                "End text"
            )

            result = clean_docx_broken_image_refs(markdown, image_dir)

            self.assertNotIn("data:image", result)
            self.assertIn("![](", result)
            self.assertIn("image1.png", result)
            self.assertIn("image2.jpg", result)
            # Original text should be preserved
            self.assertIn("Some text", result)
            self.assertIn("More text", result)
            self.assertIn("End text", result)

    def test_appends_extra_images_when_more_extracted_than_placeholders(self):
        """Extra extracted images beyond placeholder count should be appended at end."""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            image_dir = temp_path / "images"
            image_dir.mkdir(parents=True)

            (image_dir / "img_a.png").write_bytes(b"a")
            (image_dir / "img_b.png").write_bytes(b"b")
            (image_dir / "img_c.png").write_bytes(b"c")

            # Only one data URI placeholder, but three images extracted
            markdown = "Start\n\n![x](data:image/png;base64...)\n\nEnd"

            result = clean_docx_broken_image_refs(markdown, image_dir)

            self.assertNotIn("data:image", result)
            self.assertIn("img_a.png", result)
            # Extra images should appear at the end
            self.assertIn("img_b.png", result)
            self.assertIn("img_c.png", result)
            self.assertIn("Start", result)
            self.assertIn("End", result)

    def test_no_data_uri_leaves_markdown_unchanged(self):
        """If there are no data: URIs, the markdown should be returned as-is."""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            image_dir = temp_path / "images"
            image_dir.mkdir(parents=True)
            (image_dir / "img.png").write_bytes(b"data")

            markdown = "Just text\n\n![](img.png)\n\nMore text"
            result = clean_docx_broken_image_refs(markdown, image_dir)

            # Should be unchanged except possibly appended unreferenced images
            self.assertIn("Just text", result)
            self.assertIn("![](img.png)", result)

    def test_no_extracted_images_leaves_markdown_unchanged(self):
        """If the image directory is empty, markdown should be unchanged."""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            image_dir = temp_path / "images"
            image_dir.mkdir(parents=True)
            # No images in the directory

            markdown = "![](data:image/png;base64...)\n\nText"
            result = clean_docx_broken_image_refs(markdown, image_dir)

            self.assertEqual(result, markdown)


if __name__ == "__main__":
    unittest.main()
